#!/usr/bin/env python3
"""
Reproduce FRAA / blocked ELISA EC50 analyses from the original workbook.

Inputs
------
- Excel workbook with the same structure as "ED50 Values.xlsx"

Outputs
-------
A folder containing:
- parsed_ec50_data.csv
- figure_1_histogram_anti_fraa_ec50.svg
- figure_1_histogram_anti_fraa_ec50_with_gaussian.svg
- figure_2_scatter_grey_subset.svg
- figure_2_scatter_with_uncertainty.svg
- patient_ec50_table.csv
- patient_ec50_table.xlsx
- detailed_legends.txt
- exploratory_* figures and source tables

Usage
-----
python fraa_reproducible_analysis.py --input "ED50 Values.xlsx" --outdir fraa_reproducible_output
"""
from __future__ import annotations

import argparse
from pathlib import Path
from typing import Dict, Iterable, List, Tuple

import numpy as np
import openpyxl
import pandas as pd
import matplotlib.pyplot as plt
from scipy.optimize import curve_fit


DILUTIONS = np.array([3, 9, 27, 81, 243, 729, 2187], dtype=float)
LOG_DILUTIONS = np.log10(DILUTIONS)


def fourpl(logx: np.ndarray, bottom: float, top: float, hill: float, logec50: float) -> np.ndarray:
    """4-parameter logistic on the log10(dilution) axis."""
    return bottom + (top - bottom) / (1.0 + 10.0 ** ((logx - logec50) * hill))


def load_workbook_data(path: str | Path) -> pd.DataFrame:
    """Parse the original workbook layout into a tidy dataframe."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Sheet1"]
    rows = list(ws.values)

    records: List[Dict] = []

    # Patients are in rows 2..83, controls in rows 86..87 (0-based indexing in list(rows))
    sample_rows = list(range(2, 84)) + [86, 87]

    for ridx in sample_rows:
        row = rows[ridx]
        sample_raw = row[0]
        if sample_raw is None:
            continue

        sample_id = str(sample_raw).replace("\xa0", " ").strip()
        group = "Control" if sample_id.lower().startswith("control") else "Patient"

        anti_ec50 = float(row[1])
        anti_y = [float(v) for v in row[2:9]]

        blocked_ec50 = float(row[10])
        blocked_y = [float(v) for v in row[11:18]]

        patient_id = None
        for token in sample_id.split():
            if token.isdigit():
                patient_id = int(token)
                break

        records.append(
            {
                "sample_id": sample_id,
                "patient_id": patient_id,
                "group": group,
                "anti_fraa_ec50": anti_ec50,
                "blocked_ec50": blocked_ec50,
                "anti_d1": anti_y[0],
                "anti_d2": anti_y[1],
                "anti_d3": anti_y[2],
                "anti_d4": anti_y[3],
                "anti_d5": anti_y[4],
                "anti_d6": anti_y[5],
                "anti_d7": anti_y[6],
                "blocked_d1": blocked_y[0],
                "blocked_d2": blocked_y[1],
                "blocked_d3": blocked_y[2],
                "blocked_d4": blocked_y[3],
                "blocked_d5": blocked_y[4],
                "blocked_d6": blocked_y[5],
                "blocked_d7": blocked_y[6],
            }
        )

    return pd.DataFrame(records)


def dilution_values(row: pd.Series, prefix: str) -> np.ndarray:
    return row[[f"{prefix}_d{i}" for i in range(1, 8)]].astype(float).to_numpy()


def approximate_logec50_se(y_values: Iterable[float], ec50_reported: float) -> Dict[str, float]:
    """
    Estimate per-sample uncertainty of the reported EC50.

    We keep log10(EC50) fixed at the workbook value, fit the other three 4PL
    parameters to the 7 OD values, compute residual RMSE, and convert RMSE in
    signal space into an approximate SE on log10(EC50) using the local slope
    of the logistic curve at the inflection point.
    """
    y = np.asarray(list(y_values), dtype=float)
    logec50 = float(np.log10(ec50_reported))

    ymin = float(np.min(y))
    ymax = float(np.max(y))
    yrange = max(ymax - ymin, 1e-9)

    def model(logx: np.ndarray, bottom: float, top: float, hill: float) -> np.ndarray:
        return fourpl(logx, bottom, top, hill, logec50)

    p0 = [ymin, ymax, 1.0]
    bounds = (
        [ymin - 0.5 * yrange, ymax - 0.5 * yrange, 0.01],
        [ymin + 0.5 * yrange, ymax + 0.5 * yrange, 20.0],
    )

    popt, _ = curve_fit(model, LOG_DILUTIONS, y, p0=p0, bounds=bounds, maxfev=100000)
    yhat = model(LOG_DILUTIONS, *popt)

    residuals = y - yhat
    df_resid = max(1, len(y) - len(popt))  # 7 points - 3 fitted params = 4 dof
    rmse = float(np.sqrt(np.sum(residuals ** 2) / df_resid))

    bottom, top, hill = [float(v) for v in popt]
    slope = abs((top - bottom) * hill * np.log(10.0) / 4.0)

    if not np.isfinite(slope) or slope <= 0:
        se_log10_ec50 = np.nan
    else:
        se_log10_ec50 = rmse / slope

    return {
        "bottom": bottom,
        "top": top,
        "hill": hill,
        "rmse": rmse,
        "se_log10_ec50": se_log10_ec50,
    }


def make_errorbars(center: float, se_log10: float) -> Tuple[float, float]:
    """Convert ±1 SE in log10(EC50) into asymmetric bars on the original EC50 scale."""
    if not np.isfinite(se_log10) or se_log10 < 0:
        return 0.0, 0.0
    lower = center - 10 ** (np.log10(center) - se_log10)
    upper = 10 ** (np.log10(center) + se_log10) - center
    return float(lower), float(upper)


def compute_derived_metrics(df: pd.DataFrame) -> pd.DataFrame:
    """Add summary and uncertainty columns used by the figures and tables."""
    out = df.copy()

    anti_max = []
    blocked_max = []
    anti_auc = []
    blocked_auc = []
    net_auc = []
    anti_over_blocked_ratio_dilutions = []
    anti_se = []
    blocked_se = []
    anti_rmse = []
    blocked_rmse = []
    xerr_low = []
    xerr_high = []
    yerr_low = []
    yerr_high = []

    for _, row in out.iterrows():
        anti_y = dilution_values(row, "anti")
        blocked_y = dilution_values(row, "blocked")

        anti_max.append(float(np.max(anti_y)))
        blocked_max.append(float(np.max(blocked_y)))
        anti_auc.append(float(np.trapz(anti_y, LOG_DILUTIONS)))
        blocked_auc.append(float(np.trapz(blocked_y, LOG_DILUTIONS)))
        net_auc.append(float(np.trapz(anti_y - blocked_y, LOG_DILUTIONS)))

        with np.errstate(divide="ignore", invalid="ignore"):
            ratios = anti_y / blocked_y
        anti_over_blocked_ratio_dilutions.append(ratios)

        anti_stats = approximate_logec50_se(anti_y, row["anti_fraa_ec50"])
        blocked_stats = approximate_logec50_se(blocked_y, row["blocked_ec50"])

        anti_se.append(anti_stats["se_log10_ec50"])
        blocked_se.append(blocked_stats["se_log10_ec50"])
        anti_rmse.append(anti_stats["rmse"])
        blocked_rmse.append(blocked_stats["rmse"])

        xl, xh = make_errorbars(row["blocked_ec50"], blocked_stats["se_log10_ec50"])
        yl, yh = make_errorbars(row["anti_fraa_ec50"], anti_stats["se_log10_ec50"])
        xerr_low.append(xl)
        xerr_high.append(xh)
        yerr_low.append(yl)
        yerr_high.append(yh)

    out["anti_max"] = anti_max
    out["blocked_max"] = blocked_max
    out["anti_auc"] = anti_auc
    out["blocked_auc"] = blocked_auc
    out["net_auc"] = net_auc
    out["anti_se_log10_ec50"] = anti_se
    out["blocked_se_log10_ec50"] = blocked_se
    out["anti_rmse"] = anti_rmse
    out["blocked_rmse"] = blocked_rmse
    out["xerr_low"] = xerr_low
    out["xerr_high"] = xerr_high
    out["yerr_low"] = yerr_low
    out["yerr_high"] = yerr_high
    out["grey_flag"] = (out["group"] == "Patient") & (out["blocked_ec50"] > out["anti_fraa_ec50"] / 2.0)
    out["blocked_gt_half_fraa"] = np.where(out["grey_flag"], "Yes", "No")
    out["log10_ratio_anti_blocked_ec50"] = np.log10(out["anti_fraa_ec50"] / out["blocked_ec50"])

    for i in range(7):
        out[f"anti_over_blocked_ratio_d{i+1}"] = [
            float(r[i]) if np.isfinite(r[i]) else np.nan for r in anti_over_blocked_ratio_dilutions
        ]

    return out


def save_basic_data_tables(df: pd.DataFrame, outdir: Path) -> None:
    df.to_csv(outdir / "parsed_ec50_data.csv", index=False)

    patient_table = (
        df[df["group"] == "Patient"][["patient_id", "anti_fraa_ec50", "blocked_ec50", "blocked_gt_half_fraa"]]
        .rename(
            columns={
                "patient_id": "Patient ID",
                "anti_fraa_ec50": "EC50 FRAA",
                "blocked_ec50": "EC50 blocked",
                "blocked_gt_half_fraa": "Blocked > FRAA/2",
            }
        )
        .sort_values("Patient ID")
        .copy()
    )

    patient_table["EC50 FRAA"] = patient_table["EC50 FRAA"].round(4)
    patient_table["EC50 blocked"] = patient_table["EC50 blocked"].round(4)

    patient_table.to_csv(outdir / "patient_ec50_table.csv", index=False)
    with pd.ExcelWriter(outdir / "patient_ec50_table.xlsx", engine="openpyxl") as writer:
        patient_table.to_excel(writer, sheet_name="Patient EC50 Table", index=False)


def configure_log_ec50_axes(ax: plt.Axes) -> None:
    ticks = [1, 3, 9, 27, 81, 243, 729]
    labels = [str(t) for t in ticks]
    ax.set_xscale("log")
    ax.set_yscale("log")
    ax.set_xticks(ticks)
    ax.set_xticklabels(labels)
    ax.set_yticks(ticks)
    ax.set_yticklabels(labels)
    ax.set_xlim(0.5, 900)
    ax.set_ylim(0.5, 900)
    ax.grid(True, alpha=0.25)


def plot_figure_1_histogram(df: pd.DataFrame, outdir: Path) -> None:
    patients = df[df["group"] == "Patient"].copy()
    controls = df[df["group"] == "Control"].copy()

    x = patients["anti_fraa_ec50"].astype(float).to_numpy()
    bins = np.array([1, 3, 9, 27, 81, 243, 729], dtype=float)

    fig, ax = plt.subplots(figsize=(8, 5.625), dpi=220)
    ax.hist(x, bins=bins, color="#c8c8c8", edgecolor="#5a5a5a", linewidth=1.0)

    colors = ["#1f77b4", "#d62728"]
    ymax = ax.get_ylim()[1]
    for i, (_, row) in enumerate(controls.iterrows(), start=1):
        val = float(row["anti_fraa_ec50"])
        color = colors[i - 1]
        ax.axvline(val, linewidth=2, color=color)
        ax.text(
            val,
            ymax * (0.95 - 0.08 * (i - 1)),
            f"{row['sample_id']}\n{val:.2f}",
            rotation=90,
            va="top",
            ha="center",
            color=color,
            bbox=dict(facecolor="white", edgecolor="none", alpha=0.75, pad=1.5),
        )

    patient_median = float(np.median(x))
    ax.axvline(patient_median, color="black", linestyle="--", linewidth=1.5)
    ax.text(
        patient_median,
        ymax * 0.82,
        f"patient median\n{patient_median:.2f}",
        rotation=90,
        va="top",
        ha="center",
        color="black",
        bbox=dict(facecolor="white", edgecolor="none", alpha=0.75, pad=1.5),
    )

    ax.axvline(3, color="gray", linestyle=":", linewidth=1.5)
    ax.set_xscale("log")
    ax.set_xlim(1, 900)
    ax.set_xticks([3, 9, 27, 81, 243, 729])
    ax.set_xticklabels(["3", "9", "27", "81", "243", "729"])
    ax.set_xlabel("EC50 (dilution factor)")
    ax.set_ylabel("Patients (count)")
    ax.set_title("Figure 1. anti-FRAA EC50 distribution")
    ax.grid(axis="y", alpha=0.3)
    fig.tight_layout()
    fig.savefig(outdir / "figure_1_histogram_anti_fraa_ec50.svg", bbox_inches="tight")
    plt.close(fig)

    logx = np.log10(x)
    mu = float(np.mean(logx))
    sigma = float(np.std(logx, ddof=1))
    n = len(logx)
    log_bins = np.log10(bins)
    binw = np.diff(log_bins)[0]
    ygrid = np.linspace(log_bins.min(), log_bins.max(), 500)
    pdf = (1 / (sigma * np.sqrt(2 * np.pi))) * np.exp(-0.5 * ((ygrid - mu) / sigma) ** 2)
    expected_counts = n * binw * pdf

    fig, ax = plt.subplots(figsize=(8, 5.625), dpi=220)
    ax.hist(x, bins=bins, color="#c8c8c8", edgecolor="#5a5a5a", linewidth=1.0)
    ax.plot(10 ** ygrid, expected_counts, linewidth=2.0, label="Gaussian fit (log10 EC50)")
    ymax = ax.get_ylim()[1]
    for i, (_, row) in enumerate(controls.iterrows(), start=1):
        val = float(row["anti_fraa_ec50"])
        color = colors[i - 1]
        ax.axvline(val, linewidth=2, color=color)
        ax.text(
            val,
            ymax * (0.95 - 0.08 * (i - 1)),
            f"{row['sample_id']}\n{val:.2f}",
            rotation=90,
            va="top",
            ha="center",
            color=color,
            bbox=dict(facecolor="white", edgecolor="none", alpha=0.75, pad=1.5),
        )
    ax.axvline(patient_median, color="black", linestyle="--", linewidth=1.5)
    ax.axvline(3, color="gray", linestyle=":", linewidth=1.5)
    ax.set_xscale("log")
    ax.set_xlim(1, 900)
    ax.set_xticks([3, 9, 27, 81, 243, 729])
    ax.set_xticklabels(["3", "9", "27", "81", "243", "729"])
    ax.set_xlabel("EC50 (dilution factor)")
    ax.set_ylabel("Patients (count)")
    ax.set_title("anti-FRAA EC50 distribution with Gaussian fit")
    ax.grid(axis="y", alpha=0.3)
    ax.legend(frameon=False, loc="upper right")
    fig.tight_layout()
    fig.savefig(outdir / "figure_1_histogram_anti_fraa_ec50_with_gaussian.svg", bbox_inches="tight")
    plt.close(fig)


def plot_figure_2_scatter(df: pd.DataFrame, outdir: Path) -> None:
    patients_grey = df[(df["group"] == "Patient") & (df["grey_flag"])].copy()
    patients_black = df[(df["group"] == "Patient") & (~df["grey_flag"])].copy()
    controls = df[df["group"] == "Control"].copy()

    fig, ax = plt.subplots(figsize=(6.5, 6.5), dpi=220)
    ax.scatter(
        patients_grey["blocked_ec50"],
        patients_grey["anti_fraa_ec50"],
        s=28,
        color="lightgrey",
        label="Patients: blocked > FRAA/2",
    )
    ax.scatter(
        patients_black["blocked_ec50"],
        patients_black["anti_fraa_ec50"],
        s=28,
        color="black",
        label="Other patients",
    )
    ax.scatter(
        controls["blocked_ec50"],
        controls["anti_fraa_ec50"],
        s=42,
        color="red",
        label="Controls",
    )

    for _, row in controls.iterrows():
        ax.annotate(
            row["sample_id"],
            (row["blocked_ec50"], row["anti_fraa_ec50"]),
            xytext=(6, 6),
            textcoords="offset points",
            color="red",
        )

    configure_log_ec50_axes(ax)
    ax.set_xlabel("EC50 blocked")
    ax.set_ylabel("EC50 FRAA")
    ax.set_title("Figure 2. FRAA EC50 versus blocked EC50")
    ax.legend(frameon=False, loc="upper left")
    fig.tight_layout()
    fig.savefig(outdir / "figure_2_scatter_grey_subset.svg", bbox_inches="tight")
    plt.close(fig)


def plot_figure_2_with_uncertainty(df: pd.DataFrame, outdir: Path) -> None:
    patients_grey = df[(df["group"] == "Patient") & (df["grey_flag"])].copy()
    patients_black = df[(df["group"] == "Patient") & (~df["grey_flag"])].copy()
    controls = df[df["group"] == "Control"].copy()

    fig, ax = plt.subplots(figsize=(6.5, 6.5), dpi=220)

    def errplot(sub: pd.DataFrame, color: str, ms: float, z: int):
        ax.errorbar(
            sub["blocked_ec50"],
            sub["anti_fraa_ec50"],
            xerr=np.vstack([sub["xerr_low"], sub["xerr_high"]]),
            yerr=np.vstack([sub["yerr_low"], sub["yerr_high"]]),
            fmt="o",
            markersize=ms,
            color=color,
            ecolor=color,
            elinewidth=0.8 if color != "red" else 1.0,
            capsize=0,
            linestyle="none",
            zorder=z,
        )

    errplot(patients_grey, "lightgrey", 4.5, 1)
    errplot(patients_black, "black", 4.5, 2)
    errplot(controls, "red", 5.5, 3)

    for _, row in controls.iterrows():
        ax.annotate(
            row["sample_id"],
            (row["blocked_ec50"], row["anti_fraa_ec50"]),
            xytext=(6, 6),
            textcoords="offset points",
            color="red",
            fontsize=9,
        )

    configure_log_ec50_axes(ax)
    ax.set_xlabel("EC50 blocked")
    ax.set_ylabel("EC50 FRAA")
    ax.set_title("Figure 2. FRAA EC50 versus blocked EC50 with uncertainty")
    fig.tight_layout()
    fig.savefig(outdir / "figure_2_scatter_with_uncertainty.svg", bbox_inches="tight")
    plt.close(fig)

    cols = [
        "sample_id",
        "group",
        "anti_fraa_ec50",
        "blocked_ec50",
        "anti_se_log10_ec50",
        "blocked_se_log10_ec50",
        "anti_rmse",
        "blocked_rmse",
        "xerr_low",
        "xerr_high",
        "yerr_low",
        "yerr_high",
        "grey_flag",
    ]
    df[cols].to_csv(outdir / "figure_2_scatter_with_uncertainty_source_data.csv", index=False)


def plot_exploratory_figures(df: pd.DataFrame, outdir: Path) -> None:
    """Optional exploratory plots corresponding to earlier intermediate analyses."""
    patients = df[df["group"] == "Patient"].copy()
    controls = df[df["group"] == "Control"].copy()

    anti_mat = np.vstack([dilution_values(r, "anti") for _, r in patients.iterrows()])
    blocked_mat = np.vstack([dilution_values(r, "blocked") for _, r in patients.iterrows()])

    anti_med = np.nanmedian(anti_mat, axis=0)
    blocked_med = np.nanmedian(blocked_mat, axis=0)

    fig, ax = plt.subplots(figsize=(7, 5), dpi=220)
    ax.plot(DILUTIONS, anti_med, marker="o", label="Patient median anti-FRAA")
    ax.plot(DILUTIONS, blocked_med, marker="o", label="Patient median blocked")
    ax.set_xscale("log")
    ax.set_xlabel("Dilution factor")
    ax.set_ylabel("Optical density")
    ax.set_title("Exploratory: cohort median curves")
    ax.grid(True, alpha=0.25)
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(outdir / "exploratory_cohort_median_curves.svg", bbox_inches="tight")
    plt.close(fig)

    anti_norm = anti_mat / anti_mat.max(axis=1, keepdims=True)
    blocked_norm = blocked_mat / blocked_mat.max(axis=1, keepdims=True)
    anti_norm_med = np.nanmedian(anti_norm, axis=0)
    blocked_norm_med = np.nanmedian(blocked_norm, axis=0)

    fig, ax = plt.subplots(figsize=(7, 5), dpi=220)
    ax.plot(DILUTIONS, anti_norm_med, marker="o", label="Patient median anti-FRAA (normalized)")
    ax.plot(DILUTIONS, blocked_norm_med, marker="o", label="Patient median blocked (normalized)")
    ax.set_xscale("log")
    ax.set_ylim(0, 1.05)
    ax.set_xlabel("Dilution factor")
    ax.set_ylabel("Normalized optical density")
    ax.set_title("Exploratory: normalized shape comparison")
    ax.grid(True, alpha=0.25)
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(outdir / "exploratory_normalized_shape_comparison.svg", bbox_inches="tight")
    plt.close(fig)

    with np.errstate(divide="ignore", invalid="ignore"):
        ratio_mat = anti_mat / blocked_mat
    ratio_med = np.nanmedian(ratio_mat, axis=0)

    fig, ax = plt.subplots(figsize=(7, 5), dpi=220)
    ax.plot(DILUTIONS, ratio_med, marker="o")
    ax.axhline(1.0, color="gray", linestyle=":")
    ax.set_xscale("log")
    ax.set_xlabel("Dilution factor")
    ax.set_ylabel("anti-FRAA / blocked ratio")
    ax.set_title("Exploratory: patient median anti/blocked ratio by dilution")
    ax.grid(True, alpha=0.25)
    fig.tight_layout()
    fig.savefig(outdir / "exploratory_ratio_vs_dilution.svg", bbox_inches="tight")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(6.5, 5.5), dpi=220)
    p = patients
    c = controls
    ax.scatter(p["log10_ratio_anti_blocked_ec50"], p["net_auc"], s=24, color="black")
    ax.scatter(c["log10_ratio_anti_blocked_ec50"], c["net_auc"], s=36, color="red")
    for _, row in c.iterrows():
        ax.annotate(
            row["sample_id"],
            (row["log10_ratio_anti_blocked_ec50"], row["net_auc"]),
            xytext=(6, 6),
            textcoords="offset points",
            color="red",
        )
    ax.axvline(0, color="gray", linestyle=":")
    ax.set_xlabel("log10(EC50 anti-FRAA / EC50 blocked)")
    ax.set_ylabel("AUC(anti-FRAA - blocked)")
    ax.set_title("Exploratory: specificity summary scatter")
    ax.grid(True, alpha=0.25)
    fig.tight_layout()
    fig.savefig(outdir / "exploratory_specificity_scatter.svg", bbox_inches="tight")
    plt.close(fig)

    patients_summary = pd.DataFrame(
        {
            "dilution_factor": DILUTIONS,
            "patient_median_anti_fraa": anti_med,
            "patient_median_blocked": blocked_med,
            "patient_median_anti_norm": anti_norm_med,
            "patient_median_blocked_norm": blocked_norm_med,
            "patient_median_ratio_anti_over_blocked": ratio_med,
        }
    )
    patients_summary.to_csv(outdir / "exploratory_cohort_medians.csv", index=False)


def write_legends(outdir: Path) -> None:
    legend_text = """Figure 1. Distribution of anti-FRAA EC50 values in the patient cohort.

This figure shows the distribution of EC50 values obtained in the anti-FRAA ELISA for the patient cohort made available for analysis. The histogram includes only patient samples. The two control samples are shown separately and labeled with their individual EC50 values so that their positions can be compared directly with the distribution of patient values.

In this assay, EC50 refers to the dilution factor at which the fitted ELISA signal reaches one-half of its dynamic range for a given sample. Operationally, EC50 was defined as the interpolated inflection point of a sigmoidal dose-response curve fitted to the optical-density values measured in the high-throughput ELISA across a series of seven serial three-fold dilutions. Accordingly, the EC50 is not read directly from one experimental point, but is derived mathematically from the fitted curve. A higher EC50 indicates that detectable signal persists at greater dilution and therefore reflects stronger apparent antibody reactivity in the sample.

For each sample, the relationship between signal and dilution was modeled using a four-parameter logistic function of the form:

y = Bottom + (Top - Bottom) / [1 + (x / EC50)^HillSlope]

where x is the dilution factor, y is the measured optical density, Top and Bottom are the upper and lower asymptotes of the fitted response curve, and HillSlope describes the steepness of the transition between the upper and lower portions of the curve.

Figure 2. Comparison of EC50 values obtained on FRAA-coated plates and blocked control plates, including per-sample uncertainty estimates.

This figure compares, for each individual sample, the EC50 obtained in the FRAA-coated ELISA with the EC50 obtained on the corresponding blocked control plate. The y-axis shows the EC50 derived from the anti-FRAA assay performed on plates coated with folate receptor-α antigen. The x-axis shows the EC50 derived from a control ELISA plate that was treated in the same way except that it was not coated with folate receptor-α. Instead, the microplate was only blocked with a nonspecific protein solution. This blocked plate is therefore used as an estimate of background or nonspecific binding.

Each point represents one sample analyzed under both conditions. Patient samples are plotted as black symbols and the two control samples are plotted as red symbols. A subset of patient samples is highlighted in light gray. These are the samples for which EC50(blocked) > 0.5 × EC50(FRAA). These light-gray samples are flagged because they may not represent convincing antigen-specific anti-FRAA reactivity.

Per-sample uncertainty was estimated from the fit of the logistic model to the underlying seven-point ELISA dilution series. For each sample and condition, the reported EC50 was kept fixed, while the remaining three logistic parameters (Bottom, Top, and HillSlope) were fitted to the seven measured optical-density values on the log10 dilution axis. Residual deviation between the observed points and the constrained fitted curve was summarized as a root mean square error (RMSE). Because seven data points were available and three parameters were fitted in this constrained step, the residual error was based on 4 residual degrees of freedom. This residual scatter in signal space was converted into an approximate uncertainty in log10(EC50) by dividing the RMSE by the absolute slope of the logistic curve at the inflection point:

|dy / dlog10(EC50)| = (Top - Bottom) × HillSlope × ln(10) / 4

The resulting approximate standard error on log10(EC50) was then back-transformed to the original EC50 scale and plotted as asymmetric horizontal and vertical error bars.
"""
    (outdir / "detailed_legends.txt").write_text(legend_text, encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Reproduce FRAA/blocked ELISA EC50 analyses from the original workbook."
    )
    parser.add_argument("--input", required=True, help="Path to the Excel workbook, e.g. 'ED50 Values.xlsx'")
    parser.add_argument("--outdir", default="fraa_reproducible_output", help="Output directory")
    args = parser.parse_args()

    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)

    df = load_workbook_data(args.input)
    df = compute_derived_metrics(df)

    save_basic_data_tables(df, outdir)
    plot_figure_1_histogram(df, outdir)
    plot_figure_2_scatter(df, outdir)
    plot_figure_2_with_uncertainty(df, outdir)
    plot_exploratory_figures(df, outdir)
    write_legends(outdir)

    print(f"Done. Outputs written to: {outdir.resolve()}")


if __name__ == "__main__":
    main()
